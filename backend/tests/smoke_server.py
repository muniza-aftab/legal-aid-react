"""Isolated browser-test server. Never uses the user's case database."""
import io
import os
import pathlib
import sys
import tempfile
import wave

root = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root))
temporary = tempfile.TemporaryDirectory(prefix="legal-aid-browser-", ignore_cleanup_errors=True)
os.environ["DATABASE_PATH"] = str(pathlib.Path(temporary.name) / "browser.db")
os.environ["ADMIN_PASSWORD"] = "browser-test-password"
os.environ["STORE_QUERY_TEXT"] = "false"
import main
from app.utils.groq_client import LegalAnalysis, DocumentAnalysis
from app.utils.chat_service import ChatResponse, SourceDoc
from app.utils.ngo_service import add_ngo
from app.utils.lawyer_service import add_lawyer
from app.utils.text_extractor import ExtractionResult
from types import SimpleNamespace
from docx import Document
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw

fixtures = root.parent / ".local" / "browser-fixtures"
fixtures.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_paragraph("SYNTHETIC CONTRACT PREVIEW")
document.add_paragraph("The response deadline in this synthetic agreement is 1 November 2099.")
document.save(fixtures / "contract.docx")
pdf = canvas.Canvas(str(fixtures / "contract.pdf"))
pdf.drawString(50, 750, "SYNTHETIC PDF CONTRACT PREVIEW")
pdf.drawString(50, 725, "This synthetic agreement requests a response within 14 days.")
pdf.save()
image = Image.new("RGB", (1000, 400), "white")
ImageDraw.Draw(image).text((30, 30), "SYNTHETIC IMAGE NOTICE: respond within 14 days.", fill="black")
image.save(fixtures / "notice.png")

ngo = add_ngo({"ngo_name": "Browser Test Legal Support", "city": "Karachi", "categories": ["Labour Law", "Family Law"], "contact_phone": "03000000000"})
add_lawyer({"full_name": "Browser Test Advocate", "bar_license": "BROWSER-TEST-1", "city": "Karachi",
            "ngo_id": ngo["ngo_id"], "expertise_main": "Labour Law", "available": True})
main.app.router.on_startup.clear()
main._rag_loaded = True
main.index_ready = lambda: True
chunk = {"id": "browser-source", "law_name": "The Payment Of Wages Act", "section": "Section 5",
         "category": "Labour Law", "text": "Synthetic source for browser testing. Wages are payable on time.", "score": 0.8}
main.retrieve = lambda *_: [chunk]

def analyze(query, chunks, language="English"):
    return LegalAnalysis(category="Labour Law", risk_level="Medium", confidence="Medium",
        summary="Browser test analysis: your wage-payment issue was reviewed.",
        laws=["The Payment Of Wages Act — Section 5: synthetic test citation."],
        rights=["Synthetic test right."], next_steps=["Synthetic next step."], sources=[chunk])

main.analyze_with_rag = analyze
main.chat_service = lambda **kwargs: ChatResponse(reply="Browser test chat: collect your available payment records.",
    sources=[SourceDoc(title=chunk["law_name"], section=chunk["section"], snippet=chunk["text"], score=0.8)],
    confidence="medium", session_id=kwargs["session_id"])
main.transcribe_audio = lambda data, language="English": ("My employer has not paid my salary for three months.", None)
main.classify_document = lambda *_: SimpleNamespace(document_type="Legal Notice", confidence=0.9, method="test_fixture", matched_rules=[])
from app.utils import vision_ocr
vision_ocr.extract_image_text = lambda *_: "SYNTHETIC IMAGE NOTICE. Please respond within 14 days. No real case information."
main.analyze_document_with_rag = lambda text, *args: DocumentAnalysis(
    summary="Browser test document analysis completed.", key_legal_points=["A synthetic notice demands a response."],
    deadlines=["Respond within the period in the document."], confidence="Low",
    urdu_explanation="یہ صرف سافٹ ویئر کی جانچ کے لیے ایک فرضی خلاصہ ہے۔",
    total_characters=len(text), analyzed_characters=len(text), segments_analyzed=1)

audio_path = str(pathlib.Path(temporary.name) / "test.wav")
with wave.open(audio_path, "wb") as audio:
    audio.setnchannels(1)
    audio.setsampwidth(2)
    audio.setframerate(16000)
    audio.writeframes(b"\0\0" * 16000)
main.generate_tts = lambda *_: audio_path
main.tts_available = lambda: True

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(main.app, host="127.0.0.1", port=8010, log_level="warning")
