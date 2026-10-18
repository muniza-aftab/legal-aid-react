"""Regression checks use a disposable database and mocked external providers."""
from __future__ import annotations
import io
import json
import os
import pathlib
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

_temporary = tempfile.TemporaryDirectory(prefix="legal-aid-tests-", ignore_cleanup_errors=True)
os.environ["DATABASE_PATH"] = str(pathlib.Path(_temporary.name) / "test.db")
os.environ["ADMIN_PASSWORD"] = "test-only-password"
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient
import main
from app.utils import auth, chat_service, groq_client
from app.utils.db import get_connection
from app.utils.grounding import validate_citations
from app.utils.lawyer_service import add_lawyer, get_lawyer
from app.utils.ngo_service import add_ngo
from app.utils.text_extractor import ExtractionResult, extract_docx

CHUNK = {"id": "audit-1", "law_name": "The Payment Of Wages Act", "section": "Section 5",
         "category": "Labour Law", "text": "Section 5. Wages shall be paid at the specified time.", "score": 0.8}


class Features(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app, raise_server_exceptions=False)

    def setUp(self):
        auth._requests.clear()
        with get_connection() as conn:
            for table in ["applications", "lawyers", "ngos", "feedback", "users_logs", "access_sessions"]:
                conn.execute(f"DELETE FROM {table}")
        login = self.client.post("/api/admin/login", json={"password": "test-only-password"})
        self.assertEqual(login.status_code, 200, login.text)
        self.admin = {"X-Admin-Token": login.json()["token"]}

    def seed(self):
        ngo = add_ngo({"ngo_name": "Synthetic NGO", "city": "Karachi", "categories": ["Family Law"]})
        lawyer = add_lawyer({"full_name": "Synthetic Lawyer", "bar_license": "TEST-0001", "city": "Karachi",
                             "expertise_main": "Family Law", "ngo_id": ngo["ngo_id"], "available": True})
        return ngo, lawyer

    def apply(self):
        response = self.client.post("/api/applications", json={
            "full_name": "Synthetic Applicant", "phone": "03000000000", "cnic": "00000-0000000-0",
            "city": "Karachi", "case_type": "Family Law", "urgency": "Medium",
            "description": "This is a synthetic case used only for software testing.", "consent": True,
        })
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["app_id"]

    def test_admin_access_and_revocation(self):
        for path in ["/api/applications", "/api/analytics/queries", "/api/analytics/summary", "/api/ngos/stats"]:
            self.assertEqual(self.client.get(path).status_code, 401, path)
            self.assertEqual(self.client.get(path, headers=self.admin).status_code, 200, path)
        self.assertEqual(self.client.post("/api/ngos", json={"name": "Unauthorized"}).status_code, 401)
        self.assertEqual(self.client.post("/api/admin/logout", headers=self.admin).status_code, 200)
        self.assertEqual(self.client.get("/api/admin/session", headers=self.admin).status_code, 401)

    def test_bad_and_expired_tokens(self):
        self.assertEqual(self.client.get("/api/applications", headers={"X-Admin-Token": "admin-session-made-up"}).status_code, 401)
        with get_connection() as conn:
            conn.execute("UPDATE access_sessions SET expires_at=0")
        self.assertEqual(self.client.get("/api/applications", headers=self.admin).status_code, 401)

    def test_chat_ownership(self):
        session = self.client.post("/api/chat/sessions").json()
        sid = session["session_id"]
        self.assertEqual(self.client.get(f"/api/chat/{sid}/history").status_code, 401)
        headers = {"X-Chat-Token": session["token"]}
        self.assertEqual(self.client.get(f"/api/chat/{sid}/history", headers=headers).status_code, 200)
        self.assertEqual(self.client.delete(f"/api/chat/{sid}", headers=headers).status_code, 200)
        self.assertEqual(self.client.get(f"/api/chat/{sid}/history", headers=headers).status_code, 401)

    def test_consent_and_application_contract(self):
        response = self.client.post("/api/applications", json={"full_name": "Synthetic Applicant", "phone": "03000000000",
            "city": "Karachi", "case_type": "Family Law", "urgency": "Medium", "description": "A" * 35})
        self.assertEqual(response.status_code, 422)
        app_id = self.apply()
        result = self.client.get(f"/api/applications/{app_id}", headers=self.admin).json()
        self.assertEqual(result["cnic"], "00000-0000000-0")
        self.assertTrue(result["consent_at"])
        self.assertEqual(self.client.get("/api/applications?city=Lahore", headers=self.admin).json()["applications"], [])
        self.assertEqual(self.client.patch(f"/api/applications/{app_id}/status", json={"status": "invalid"}, headers=self.admin).status_code, 422)

    def test_workload_and_terminal_status(self):
        _, lawyer = self.seed()
        app_id = self.apply()
        self.assertEqual(get_lawyer(lawyer["lawyer_id"])["active_cases"], 1)
        for _ in range(2):
            response = self.client.patch(f"/api/applications/{app_id}/lawyer", json={"lawyer_id": lawyer["lawyer_id"]}, headers=self.admin)
            self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(get_lawyer(lawyer["lawyer_id"])["active_cases"], 1)
        self.assertEqual(self.client.patch(f"/api/applications/{app_id}/status", json={"status": "resolved"}, headers=self.admin).status_code, 200)
        self.assertEqual(get_lawyer(lawyer["lawyer_id"])["active_cases"], 0)
        self.assertNotEqual(self.client.post(f"/api/applications/{app_id}/reroute", headers=self.admin).status_code, 200)
        self.assertEqual(self.client.get(f"/api/applications/{app_id}", headers=self.admin).json()["status"], "resolved")

    def test_feedback_and_validation(self):
        self.assertEqual(self.client.post("/api/feedback", json={"session_id": "test", "rating": 5}).status_code, 200)
        self.assertEqual(self.client.post("/api/feedback", json={"rating": 10}).status_code, 422)
        self.assertEqual(self.client.post("/api/analyze", json={"query": "A" * 2001}).status_code, 422)
        self.assertEqual(self.client.post("/api/analyze", json={"query": "A" * 30, "language": "unknown"}).status_code, 422)

    def test_private_analytics_default(self):
        from app.utils.db import log_query
        with patch.dict(os.environ, {"STORE_QUERY_TEXT": "false"}):
            row_id = log_query("English", "Synthetic confidential question", "Labour Law", "Medium", "Synthetic confidential summary", city="Karachi")
        with get_connection() as conn:
            row = conn.execute("SELECT * FROM users_logs WHERE id=?", (row_id,)).fetchone()
        self.assertEqual((row["query_text"], row["ai_summary"]), ("", ""))
        self.assertEqual((row["city"], row["category"]), ("Karachi", "Labour Law"))

    def test_historical_assignments_survive_registry_deletion_attempts(self):
        ngo, lawyer = self.seed()
        app_id = self.apply()
        self.client.patch(f"/api/applications/{app_id}/status", json={"status": "resolved"}, headers=self.admin)
        self.assertEqual(self.client.delete(f'/api/lawyers/{lawyer["lawyer_id"]}', headers=self.admin).status_code, 409)
        self.assertEqual(self.client.delete(f'/api/ngos/{ngo["ngo_id"]}', headers=self.admin).status_code, 409)
        case = self.client.get(f"/api/applications/{app_id}", headers=self.admin).json()
        self.assertEqual((case["assigned_lawyer_id"], case["assigned_ngo_id"]), (lawyer["lawyer_id"], ngo["ngo_id"]))

    def test_inactive_ngos_and_lawyer_specialization(self):
        ngo, _ = self.seed()
        self.client.patch(f'/api/ngos/{ngo["ngo_id"]}/deactivate', headers=self.admin)
        self.assertEqual(self.client.get("/api/ngos").json()["total"], 0)
        self.assertEqual(self.client.get("/api/ngos", headers=self.admin).json()["total"], 1)
        result = self.client.post("/api/lawyers", json={"full_name": "Specialist Lawyer", "bar_license": "TEST-0002",
            "city": "Karachi", "expertise_main": "Family Law", "expertise_sub": ["Divorce"]}, headers=self.admin)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["expertise_sub"], ["Divorce"])

    def test_risk_and_negation(self):
        self.assertNotEqual(groq_client.stabilize_risk("Medium", "I need help with a skill training employment contract.", "Labour Law")[0], "Critical")
        self.assertEqual(groq_client.stabilize_risk("High", "What is the law about murder in Pakistan?", "Criminal Law")[0], "Low")
        self.assertEqual(groq_client.stabilize_risk("High", "How do I report harassment by my employer?", "Harassment")[0], "High")
        facts = chat_service._extract_case_facts("I have no written contract and no witnesses.")
        self.assertEqual(facts, ["User stated: I have no written contract and no witnesses."])

    def test_citation_validation(self):
        good = "The Payment Of Wages Act — Section 5: wages provision."
        self.assertEqual(validate_citations([good], [1], [CHUNK])[0], [good])
        self.assertEqual(validate_citations(["Invented Act — Section 999: fabricated."], [1], [CHUNK])[0], [])
        self.assertEqual(validate_citations(["The Payment Of Wages Act — Section 999: fabricated."], [1], [CHUNK])[0], [])
        with patch.object(groq_client, "call_groq_with_rotation") as provider:
            result = groq_client.analyze_with_rag("My employer has not paid my wages for three months.", [])
            provider.assert_not_called()
            self.assertEqual((result.laws, result.confidence), ([], "Low"))

    def test_document_segments_cover_ending(self):
        text = "A" * 11000 + "DEADLINE_AT_END"
        observed = []
        def analyze(segment, *_):
            observed.append(segment)
            return groq_client.DocumentAnalysis(summary="Synthetic section", deadlines=["End deadline"] if "DEADLINE_AT_END" in segment else [], confidence="Low")
        with patch.object(groq_client, "_analyze_document_chunk", side_effect=analyze):
            result = groq_client.analyze_document_with_rag(text, "Legal Notice", "English")
        self.assertEqual("".join(observed), text)
        self.assertEqual(result.analyzed_characters, len(text))
        self.assertFalse(result.truncated)
        self.assertIn("End deadline", result.deadlines)

    def test_txt_upload_and_invalid_file(self):
        cls = SimpleNamespace(document_type="Legal Notice", confidence=0.9, method="rule_based", matched_rules=[])
        with patch.object(main, "classify_document", return_value=cls), patch.object(main, "analyze_document_with_rag", return_value=groq_client.DocumentAnalysis(summary="Synthetic analysis")):
            response = self.client.post("/api/analyze-document", files={"file": ("sample.txt", b"This is a synthetic legal notice for a software feature test.", "text/plain")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("synthetic", response.json()["extracted_text"])
        bad = self.client.post("/api/analyze-document", files={"file": ("bad.pdf", b"Not a PDF file", "application/pdf")})
        self.assertEqual(bad.status_code, 422)

    def test_docx_includes_tables(self):
        from docx import Document
        path = str(pathlib.Path(_temporary.name) / "table.docx")
        doc = Document()
        doc.add_paragraph("Synthetic agreement")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "A deadline inside a table"
        doc.save(path)
        self.assertIn("deadline inside a table", extract_docx(path).text)

    def test_pdf_export_and_cors(self):
        response = self.client.post("/api/pdf", json={"summary": "Synthetic test report", "category": "Labour Law", "laws": [], "rights": [], "next_steps": []})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.content.startswith(b"%PDF"))
        cors = self.client.options("/api/config", headers={"Origin": "https://untrusted.invalid", "Access-Control-Request-Method": "GET"})
        self.assertEqual(cors.status_code, 400)


if __name__ == "__main__":
    unittest.main()
